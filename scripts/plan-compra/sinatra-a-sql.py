#!/usr/bin/env python3
"""Sinatra → SQL: lo vendido cada día por tienda, para `retail.ventas_historicas` (ADR-0372, C2 · Felipe 2026-10-12).

Lee los Excel «Sinatra» de Finanzas (la carpeta SINATRAS de Drive, sincronizada en la Mac) y escribe un archivo SQL que se pega UNA vez
en la base. Los datos nunca entran al repo (es público): el SQL va a la ruta que le pases, fuera del repo.

    python3 scripts/plan-compra/sinatra-a-sql.py "<carpeta SINATRAS>" <salida.sql>

Requiere `openpyxl` (pip install --user openpyxl).

QUÉ LEE (y por qué esos archivos; lo decidido con Felipe el 2026-10-12):
  · De cada año, el archivo final (no las «Copia de…», «revisar…» ni «Con margen ERRADO»). Si un archivo cambia de nombre, se
    corrige la tabla FUENTES de abajo, no se adivina.
  · Hasta 2023 hay un solo «Diario»: es Trujillo (Arequipa abre en 2024 con su propio Diario; Lima, en 2026).
  · Cada archivo aporta SOLO los días de su año: «Sinatra 2017 - inicios 2018» trae enero de 2018, que ya está en el de 2018, y un día
    contado dos veces inflaría ese enero.
  · Fuera: KIDS y «Sodalita 2023» (la línea de niños: Sodalita es el mismo diario que KIDS con S/ 300 más por día), «CAYLA PET 2022»
    (una copia de 2022) y Lima 2026 (su hoja «Ingresos» tiene 20 registros sueltos en siete meses, no un diario: un multiplicador
    calculado con eso mentiría).

LA TIENDA se reconoce por su ciudad, no por su nombre exacto: en producción se llama «Tienda TRU» y en una base local «Tienda Trujillo».
Si una base no tiene esa tienda (una local sin Arequipa), sus días no se cargan y el SQL lo avisa con un NOTICE, sin fallar.
"""

import datetime
import os
import sys
import warnings

import openpyxl

warnings.filterwarnings("ignore")

# (año, carpeta, archivo, hoja, tienda)
FUENTES = [
    (2017, "Sinatra 2017", "Sinatra 2017 - inicios 2018.xlsm", "Diario", "TRU"),
    (2018, "Sinatra 2018", "Sinatra 2018.xlsm", "Diario", "TRU"),
    (2019, "Sinatra 2019", "Sinatra 2019 (1).xlsm", "Diario", "TRU"),
    (2020, "Sinatra 2020", "Sinatra 2020 ORIGINAL Final.xlsm", "Diario", "TRU"),
    (2021, "Sinatra 2021", "Sinatra Final 2021 .xlsm", "Diario", "TRU"),
    (2022, "Sinatra 2022", "SINATRA 2022 .xlsm", "Diario", "TRU"),
    (2023, "Sinatra 2023", "SINATRA 2023 .xlsm", "Diario", "TRU"),
    (2024, "Sinatra 2024", "SINATRA 2024 .xlsm", "Diario TRU", "TRU"),
    (2024, "Sinatra 2024", "SINATRA 2024 .xlsm", "Diario AQP", "AQP"),
    (2025, "Sinatra 2025", "SINATRA 2025.xlsm", "Diario TRU", "TRU"),
    (2025, "Sinatra 2025", "SINATRA 2025.xlsm", "Diario AQP", "AQP"),
    (2026, "Sinatra 2026", "SINATRA 2026 TRU.xlsm", "Diario TRU", "TRU"),
    (2026, "Sinatra 2026", "SINATRA 2026 AQP.xlsm", "Diario AQP", "AQP"),
]

# Cómo se encuentra cada tienda en la base: por su ciudad o su sigla, como palabra suelta (\m…\M), sin importar mayúsculas.
PATRON_TIENDA = {"TRU": r"\m(tru|trujillo)\M", "AQP": r"\m(aqp|arequipa)\M"}


def leer_diario(ruta, hoja, anio):
    """{fecha: soles} de la hoja: la fila de títulos trae «Fecha» y «Ventas…»; un día con varias filas se suma; solo ventas > 0."""
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    if hoja not in wb.sheetnames:
        raise SystemExit(f"«{os.path.basename(ruta)}» no tiene la hoja «{hoja}»: corrige FUENTES.")
    dias, cab = {}, None
    for fila in wb[hoja].iter_rows(values_only=True):
        if cab is None:
            celdas = [str(x).strip() if x else "" for x in fila[:15]]
            if "Fecha" in celdas and any(c.startswith("Ventas") for c in celdas):
                cab = (celdas.index("Fecha"), next(i for i, c in enumerate(celdas) if c.startswith("Ventas")))
            continue
        fecha = fila[cab[0]] if len(fila) > cab[0] else None
        soles = fila[cab[1]] if len(fila) > cab[1] else None
        if isinstance(fecha, datetime.datetime) and fecha.year == anio and isinstance(soles, (int, float)) and soles > 0:
            dias[fecha.date()] = dias.get(fecha.date(), 0.0) + float(soles)
    if cab is None:
        raise SystemExit(f"«{os.path.basename(ruta)} · {hoja}» no tiene la fila de títulos «Fecha / Ventas».")
    return dias


def main():
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    carpeta, salida = sys.argv[1], sys.argv[2]
    if os.path.abspath(salida).startswith(os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))):
        raise SystemExit("La salida no puede quedar dentro del repo (es público): elige una ruta afuera.")
    filas, resumen = [], []
    for anio, sub, archivo, hoja, tienda in FUENTES:
        dias = leer_diario(os.path.join(carpeta, sub, archivo), hoja, anio)
        fuente = f"{archivo.strip()} · {hoja}".replace("'", "''")
        filas += [f"('{tienda}', '{d.isoformat()}', {s:.2f}, '{fuente}')" for d, s in sorted(dias.items())]
        resumen.append(f"--   {anio} {tienda}: {len(dias)} días, S/ {sum(dias.values()):,.2f} ({fuente})")
    casos = "\n".join(f"      when '{t}' then (select array_agg(u.id) from retail.ubicaciones u where u.tipo = 'tienda' and u.nombre ~* '{p}')"
                      for t, p in PATRON_TIENDA.items())
    sql = f"""-- Carga de Sinatra a retail.ventas_historicas (ADR-0372, C2). Generado el {datetime.date.today()} por scripts/plan-compra/sinatra-a-sql.py.
-- NO va al repo. Se pega una vez (o se vuelve a pegar tras corregir un Excel: pisa los mismos días).
{chr(10).join(resumen)}
set search_path = retail, public, extensions;
do $carga$
declare
  r record;
  v_ids uuid[];
  v_cargadas integer := 0;
  v_saltadas integer := 0;
begin
  for r in select * from (values
{(',' + chr(10)).join('    ' + f for f in filas)}
  ) as t(tienda, fecha, soles, fuente) loop
    v_ids := case r.tienda
{casos}
    end;
    if coalesce(array_length(v_ids, 1), 0) <> 1 then
      v_saltadas := v_saltadas + 1;
      continue;
    end if;
    insert into retail.ventas_historicas (ubicacion_id, fecha, soles, fuente)
    values (v_ids[1], r.fecha::date, r.soles, r.fuente)
    on conflict (ubicacion_id, fecha) do update set soles = excluded.soles, fuente = excluded.fuente, cargado_en = now();
    v_cargadas := v_cargadas + 1;
  end loop;
  raise notice 'Sinatra: % días cargados; % saltados (su tienda no existe en esta base o hay más de una que calza).', v_cargadas, v_saltadas;
end
$carga$;
"""
    with open(salida, "w", encoding="utf-8") as f:
        f.write(sql)
    print(f"{len(filas)} días escritos en {salida}")
    print("\n".join(r[4:] for r in resumen))


if __name__ == "__main__":
    main()
